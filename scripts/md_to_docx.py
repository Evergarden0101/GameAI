#!/usr/bin/env python3
"""
Build a polished, navigable .docx from the Game AI book Markdown.

A small Markdown -> Word converter built on `python-docx`. It emits real Word
heading styles (so the document has a working Navigation pane and an auto-updating
Table of Contents field), styled tables, shaded code blocks, block quotes, and a
cover page.

Usage:
    pip install python-docx
    python scripts/md_to_docx.py docs/game-ai-book.en.md

Optional args: <md_path> <title> <subtitle> <byline> <out.docx>
"""
import re
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor

ACCENT = RGBColor(0x2B, 0x5C, 0x8A)
DARK = RGBColor(0x12, 0x23, 0x3B)
CODE_BG = "F1F4F7"
HEAD_BG = "2B5C8A"


# --------------------------------------------------------------------------- #
# low-level helpers
# --------------------------------------------------------------------------- #
# OOXML enforces child ordering inside <w:pPr> and <w:tcPr>; insert accordingly.
PPR_ORDER = ["pStyle", "keepNext", "keepLines", "pageBreakBefore", "framePr",
             "widowControl", "numPr", "suppressLineNumbers", "pBdr", "shd",
             "tabs", "suppressAutoHyphens", "kinsoku", "wordWrap", "overflowPunct",
             "topLinePunct", "autoSpaceDE", "autoSpaceDN", "bidi", "adjustRightInd",
             "snapToGrid", "spacing", "ind", "contextualSpacing", "mirrorIndents",
             "suppressOverlap", "jc", "textDirection", "textAlignment",
             "textboxTightWrap", "outlineLvl", "divId", "cnfStyle", "rPr",
             "sectPr", "pPrChange"]
TCPR_ORDER = ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders",
              "shd", "noWrap", "tcMar", "textDirection", "tcFit", "vAlign",
              "hideMark", "tcPrChange"]


def _insert_ordered(parent, child, order):
    tag = child.tag.split("}")[-1]
    rank = order.index(tag)
    for existing in parent:
        etag = existing.tag.split("}")[-1]
        if etag in order and order.index(etag) > rank:
            existing.addprevious(child)
            return
    parent.append(child)


def shade(cell_or_para, hex_fill):
    if hasattr(cell_or_para, "_tc"):
        parent, order = cell_or_para._tc.get_or_add_tcPr(), TCPR_ORDER
    else:
        parent, order = cell_or_para._p.get_or_add_pPr(), PPR_ORDER
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    _insert_ordered(parent, shd, order)


def left_border(paragraph, hex_color="2B5C8A"):
    pPr = paragraph._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), hex_color)
    pbdr.append(left)
    _insert_ordered(pPr, pbdr, PPR_ORDER)


INLINE_RE = re.compile(
    r"(`[^`]+`)"                      # code
    r"|(\[[^\]]+\]\([^)]+\))"         # link
    r"|(\*\*[^*]+\*\*)"               # bold
    r"|(\*[^*]+\*)"                   # italic
)


def add_inline(paragraph, text):
    """Split `text` into runs honoring **bold**, *italic*, `code`, [link](url)."""
    pos = 0
    for m in INLINE_RE.finditer(text):
        if m.start() > pos:
            paragraph.add_run(text[pos:m.start()])
        tok = m.group(0)
        if tok.startswith("`"):
            r = paragraph.add_run(tok[1:-1])
            r.font.name = "Consolas"
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0x9B, 0x25, 0x53)
        elif tok.startswith("["):
            label = tok[1:tok.index("]")]
            r = paragraph.add_run(label)
            r.font.color.rgb = ACCENT
            r.underline = True
        elif tok.startswith("**"):
            paragraph.add_run(tok[2:-2]).bold = True
        else:
            paragraph.add_run(tok[1:-1]).italic = True
        pos = m.end()
    if pos < len(text):
        paragraph.add_run(text[pos:])


def add_toc(doc):
    p = doc.add_paragraph()
    run = p.add_run()
    for kind, txt in (("begin", None), ("instr", ' TOC \\o "1-3" \\h \\z \\u '),
                      ("separate", None), ("text", "Update this field in Word / "
                                           "LibreOffice (right-click → Update) to "
                                           "build the table of contents from the "
                                           "headings."), ("end", None)):
        el = OxmlElement("w:fldChar") if kind in ("begin", "separate", "end") else \
            OxmlElement("w:instrText" if kind == "instr" else "w:t")
        if kind in ("begin", "separate", "end"):
            el.set(qn("w:fldCharType"), kind)
        else:
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


# --------------------------------------------------------------------------- #
# CJK: set an East-Asian font on every run + the core styles so Chinese renders
# with a proper face instead of a substitution.
# --------------------------------------------------------------------------- #
def apply_cjk(doc, font="Microsoft YaHei"):
    def set_ea(rpr):
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)                 # rFonts must be first in rPr
        rfonts.set(qn("w:eastAsia"), font)

    def all_runs(paras):
        for p in paras:
            for r in p.runs:
                yield r

    for r in all_runs(doc.paragraphs):
        set_ea(r._r.get_or_add_rPr())
    for t in doc.tables:
        for row in t.rows:
            for cell in row.cells:
                for r in all_runs(cell.paragraphs):
                    set_ea(r._r.get_or_add_rPr())
    for sname in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Heading 4"):
        try:
            set_ea(doc.styles[sname].element.get_or_add_rPr())
        except KeyError:
            pass


# --------------------------------------------------------------------------- #
# markdown -> docx
# --------------------------------------------------------------------------- #
def build(md_path, title, subtitle, byline, out_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().split("\n")

    doc = Document()

    # ensure the template's <w:zoom> carries a percent (schema requires it)
    settings = doc.settings.element
    zoom = settings.find(qn("w:zoom"))
    if zoom is None:
        zoom = OxmlElement("w:zoom")
        settings.insert(0, zoom)
    zoom.set(qn("w:percent"), "100")

    # page setup: US Letter, comfortable margins
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    for m in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(sec, m, Inches(0.9))

    # base font
    normal = doc.styles["Normal"]
    normal.font.name = "Georgia"
    normal.font.size = Pt(10.5)

    # ---- cover page ----
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = 1
    r = t.add_run(title)
    r.bold = True
    r.font.size = Pt(30)
    r.font.color.rgb = DARK
    r.font.name = "Arial"
    s = doc.add_paragraph()
    s.alignment = 1
    rs = s.add_run(subtitle)
    rs.font.size = Pt(14)
    rs.font.color.rgb = ACCENT
    rs.font.name = "Arial"
    doc.add_paragraph()
    b = doc.add_paragraph()
    b.alignment = 1
    rb = b.add_run(byline)
    rb.italic = True
    rb.font.size = Pt(11)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ---- table of contents ----
    h = doc.add_paragraph()
    hr = h.add_run("Contents")
    hr.bold = True
    hr.font.size = Pt(18)
    hr.font.name = "Arial"
    hr.font.color.rgb = DARK
    add_toc(doc)
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    dropped_title = dropped_sub = False
    i = 0
    while i < len(lines):
        line = lines[i]

        # fenced code
        if line.startswith("```"):
            i += 1
            code = []
            while i < len(lines) and not lines[i].startswith("```"):
                code.append(lines[i])
                i += 1
            i += 1
            p = doc.add_paragraph()
            shade(p, CODE_BG)
            left_border(p)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(4)
            for j, cl in enumerate(code):
                run = p.add_run(cl)
                run.font.name = "Consolas"
                run.font.size = Pt(8.5)
                if j < len(code) - 1:
                    run.add_break(WD_BREAK.LINE)
            continue

        # table
        if line.startswith("|") and i + 1 < len(lines) and re.match(r"^\|[\s:|-]+\|$", lines[i + 1]):
            header = [c.strip() for c in line.strip("|").split("|")]
            i += 2
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                rows.append([c.strip() for c in lines[i].strip("|").split("|")])
                i += 1
            table = doc.add_table(rows=1, cols=len(header))
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for k, htext in enumerate(header):
                cell = table.rows[0].cells[k]
                cell.paragraphs[0].text = ""
                run = cell.paragraphs[0].add_run(htext)
                run.bold = True
                run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                run.font.size = Pt(9.5)
                run.font.name = "Arial"
                shade(cell, HEAD_BG)
            for r_i, rowdata in enumerate(rows):
                cells = table.add_row().cells
                for k, ctext in enumerate(rowdata[:len(header)]):
                    para = cells[k].paragraphs[0]
                    add_inline(para, ctext)
                    for run in para.runs:
                        run.font.size = Pt(9)
                    if r_i % 2 == 1:
                        shade(cells[k], "F3F6F9")
            doc.add_paragraph()
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", line)
        if m:
            lvl, text = len(m.group(1)), m.group(2)
            if lvl == 1 and not dropped_title:
                dropped_title = True
                i += 1
                continue
            if lvl == 3 and not dropped_sub:
                dropped_sub = True
                i += 1
                continue
            hp = doc.add_heading(level=min(lvl, 4))
            add_inline(hp, text)
            if lvl == 1:
                hp.paragraph_format.page_break_before = True
            i += 1
            continue

        # hr
        if re.match(r"^---+$", line):
            i += 1
            continue

        # blockquote
        if line.startswith(">"):
            q = []
            while i < len(lines) and lines[i].startswith(">"):
                q.append(lines[i].lstrip(">").strip())
                i += 1
            p = doc.add_paragraph()
            shade(p, "EEF4FB")
            left_border(p)
            add_inline(p, " ".join(q))
            continue

        # unordered list
        if re.match(r"^\s*[-*]\s+", line):
            while i < len(lines) and re.match(r"^\s*[-*]\s+", lines[i]):
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, re.sub(r"^\s*[-*]\s+", "", lines[i]))
                i += 1
            continue

        # ordered list
        if re.match(r"^\s*\d+\.\s+", line):
            while i < len(lines) and re.match(r"^\s*\d+\.\s+", lines[i]):
                p = doc.add_paragraph(style="List Number")
                add_inline(p, re.sub(r"^\s*\d+\.\s+", "", lines[i]))
                i += 1
            continue

        if not line.strip():
            i += 1
            continue

        # paragraph
        para = [line]
        i += 1
        while i < len(lines) and lines[i].strip() and not re.match(
                r"^(#{1,6}\s|```|\||>|\s*[-*]\s|\s*\d+\.\s|---+$)", lines[i]):
            para.append(lines[i])
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        add_inline(p, " ".join(para))

    if ".zh." in md_path or "zh" in out_path.rsplit("/", 1)[-1]:
        apply_cjk(doc)

    doc.save(out_path)
    print("wrote", out_path)


if __name__ == "__main__":
    md = sys.argv[1] if len(sys.argv) > 1 else "docs/game-ai-book.en.md"
    title = sys.argv[2] if len(sys.argv) > 2 else "Game AI: From Pac-Man to GT Sophy"
    subtitle = sys.argv[3] if len(sys.argv) > 3 else \
        "An Illustrated, Hands-On Guide to the Techniques Behind Game Characters"
    byline = sys.argv[4] if len(sys.argv) > 4 else \
        "A companion textbook to the GameAI repository"
    out = sys.argv[5] if len(sys.argv) > 5 else md.rsplit(".", 2)[0] + ".docx"
    build(md, title, subtitle, byline, out)
